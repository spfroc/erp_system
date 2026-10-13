from pathlib import Path

from openpyxl import Workbook

from scripts.import_init_data import load_directory


def write_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "客户信息"
    sheet.append(["公司名称", "联系人", "联系电话"])
    sheet.append(["测试客户有限公司", "张三", "13000000000"])

    sheet = workbook.create_sheet("渠道信息")
    sheet.append(["公司名称", "联系人", "收款信息"])
    sheet.append(["测试供应商有限公司", "李四", "银行账号"])

    sheet = workbook.create_sheet("价格体系")
    sheet.append(["商品名称", "型号", "渠道", "成本单价", "售价"])
    sheet.append(["空调", "KFR-35", "测试供应商有限公司", "2000", "2600"])

    sheet = workbook.create_sheet("订单管理")
    sheet.append(["自动编号", "客户信息", "使用单位", "商品名称", "金额", "状态", "订单来源平台"])
    sheet.append(["SO-001", "测试客户有限公司", "测试学校", "空调", "2600", "进行中", "测试平台"])

    workbook.save(path)


def test_load_directory_builds_normalized_payload(tmp_path):
    write_workbook(tmp_path / "legacy.xlsx")

    payload = load_directory(tmp_path)

    assert payload["import_report"]["counts"]["organizations"] == 3
    assert payload["import_report"]["counts"]["products"] == 1
    assert payload["import_report"]["counts"]["channel_prices"] == 1
    assert payload["import_report"]["counts"]["projects"] == 1
    assert payload["import_report"]["counts"]["orders"] == 1

    customer = next(item for item in payload["organizations"] if item["name"] == "测试客户有限公司")
    assert "customer" in customer["roles"]

    supplier = next(item for item in payload["organizations"] if item["name"] == "测试供应商有限公司")
    assert "supplier" in supplier["roles"]

