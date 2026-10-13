from __future__ import annotations

import logging
import os
import re
import shutil
import subprocess
import tempfile
from copy import deepcopy
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZipInfo

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

logger = logging.getLogger(__name__)

TEMPLATE_DIR = Path(__file__).resolve().parent / "contract_templates"
PLATFORM_TEMPLATE_FILES = {
    "泉E采": TEMPLATE_DIR / "qec.docx",
    "青慧采": TEMPLATE_DIR / "qhc.docx",
}
PLATFORM_DEFAULTS = {
    "泉E采": {"shopCategories": ["服务类"], "listingLimit": "300", "amount": "3000"},
    "青慧采": {"shopCategories": ["产品类"], "listingLimit": "不限数量", "amount": "4000"},
}
SHOP_CATEGORY_OPTIONS = (
    ("服务类", "（政企专用）服务/工程类"),
    ("产品类", "货物类"),
)
SIGNER_BLANK = "\u3000" * 8
WORD_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_CONTENT_TYPE = "application/pdf"
_SOFFICE_CANDIDATES = (
    Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
    Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
    Path("/usr/bin/soffice"),
    Path("/usr/bin/libreoffice"),
    Path("/usr/lib/libreoffice/program/soffice"),
)
TOKEN_RE = re.compile(r"\{\{[A-Z0-9_]+\}\}")
SIGNER_TOKENS = {"{{PARTY_A_SIGNER}}", "{{PARTY_B_SIGNER}}"}
# 开户行/账号写死在各平台合同模板中；草稿仅回填展示用，生成时不再改写模板。
TEMPLATE_BANK_BY_PLATFORM = {
    "泉E采": "中信银行济南市中支行",
    "青慧采": "广发银行股份有限公司济南分行营业部",
}
TEMPLATE_ACCOUNT_BY_PLATFORM = {
    "泉E采": "8112501012401497344",
    "青慧采": "9550880219307000169",
}
TEMPLATE_BANK_NAME = TEMPLATE_BANK_BY_PLATFORM["泉E采"]
TEMPLATE_ACCOUNT_NUMBER = TEMPLATE_ACCOUNT_BY_PLATFORM["泉E采"]
FOOTER_LEFT_TEXT = "链接平台 链接厂商"
FOOTER_RIGHT_TEXT = "链接数据 创造价值"
ACCOUNT_FIELDS = ("accountName",)  # 开户行/账号固定写在模板中，仅户名可改
REQUIRED_FIELDS = (
    "partyAName",
    "partyAAddress",
    "partyAPhone",
    "partyBName",
    "partyBAddress",
    "partyBPhone",
    "shopName",
    "shopCategory",
    "listingLimit",
    "startDate",
    "endDate",
    "signDate",
    *ACCOUNT_FIELDS,
)
CN_DIGITS = "零壹贰叁肆伍陆柒捌玖"
_COMPANY_SUFFIXES = ("有限责任公司", "有限公司")


def normalize_contract_platform(platform: str | None) -> str:
    text = (platform or "").strip()
    if not text:
        return ""
    folded = text.replace("Ｅ", "E").replace("ｅ", "e").replace("E", "e")
    if folded == "泉e采":
        return "泉E采"
    if text == "青慧采":
        return "青慧采"
    return ""



def template_bank_account(platform: str | None) -> tuple[str, str]:
    key = normalize_contract_platform(platform)
    return (
        TEMPLATE_BANK_BY_PLATFORM.get(key, TEMPLATE_BANK_NAME),
        TEMPLATE_ACCOUNT_BY_PLATFORM.get(key, TEMPLATE_ACCOUNT_NUMBER),
    )


def has_service_contract_template(platform: str | None) -> bool:
    key = normalize_contract_platform(platform)
    path = PLATFORM_TEMPLATE_FILES.get(key)
    return bool(path and path.is_file())


def contract_needs_collection_account(platform: str | None) -> bool:
    return bool(normalize_contract_platform(platform))


def contract_needs_payment_qr(platform: str | None) -> bool:
    return normalize_contract_platform(platform) == "泉E采"


def shop_name_from_company(name: str) -> str:
    text = (name or "").strip()
    for suffix in _COMPANY_SUFFIXES:
        if text.endswith(suffix):
            text = text[: -len(suffix)]
            break
    text = text.strip()
    if text and not text.endswith("店"):
        return f"{text}店"
    return text


def shop_display_name(name: str) -> str:
    text = (name or "").strip()
    if text and not text.endswith("店"):
        return f"{text}店"
    return text


def format_shop_category(draft: dict) -> str:
    raw = draft.get("shopCategories")
    values: list[str] = []
    if isinstance(raw, list):
        values = [str(item).strip() for item in raw if str(item).strip()]
    elif str(draft.get("shopCategory") or "").strip():
        text = str(draft.get("shopCategory") or "").strip()
        labels = {label for _, label in SHOP_CATEGORY_OPTIONS}
        if text in labels or "/" in text:
            return text
        values = [part.strip() for part in re.split(r"[/,，]", text) if part.strip()]
    selected = set(values)
    return "/".join(label for value, label in SHOP_CATEGORY_OPTIONS if value in selected)


def amount_to_chinese(value) -> str:
    amount = Decimal(str(value or 0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    negative = amount < 0
    amount = abs(amount)
    yuan = int(amount)
    cents = int((amount - Decimal(yuan)) * 100)
    jiao, fen = divmod(cents, 10)
    yuan_part = f"{_integer_to_chinese(yuan)}元" if yuan else ""
    if jiao == 0 and fen == 0:
        text = f"{yuan_part or '零元'}整"
        return f"负{text}" if negative else text
    parts = [yuan_part or "零元"]
    if jiao:
        parts.append(f"{CN_DIGITS[jiao]}角")
    elif fen:
        parts.append("零")
    if fen:
        parts.append(f"{CN_DIGITS[fen]}分")
    text = "".join(parts)
    return f"负{text}" if negative else text


def format_amount(value) -> str:
    amount = Decimal(str(value or 0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if amount == amount.to_integral_value():
        return str(int(amount))
    return f"{amount:.2f}"


def parse_date_parts(value: str) -> tuple[str, str, str]:
    match = re.search(r"(\d{4})\D+(\d{1,2})\D+(\d{1,2})", (value or "").strip())
    if not match:
        return "", "", ""
    year, month, day = match.groups()
    return year, str(int(month)), str(int(day))


def normalize_contract_format(value: str | None) -> str:
    text = str(value or "word").strip().lower()
    if text in {"", "word", "docx", "doc"}:
        return "word"
    if text == "pdf":
        return "pdf"
    return ""


def fill_contract_document(platform: str, values: dict[str, str]) -> bytes:
    key = normalize_contract_platform(platform)
    path = PLATFORM_TEMPLATE_FILES.get(key)
    if not path or not path.is_file():
        raise ValueError("当前入驻平台没有可用的合同模板。")
    document = Document(str(path))
    mapping = {f"{{{{{name}}}}}": str(values.get(name) or "") for name in (
        "PARTY_A_NAME",
        "PARTY_B_NAME",
        "SHOP_NAME",
        "SHOP_CATEGORY",
        "LISTING_LIMIT",
        "START_YEAR",
        "START_MONTH",
        "START_DAY",
        "END_YEAR",
        "END_MONTH",
        "END_DAY",
        "AMOUNT",
        "AMOUNT_CN",
        "ACCOUNT_NAME",
        "PARTY_A_ADDRESS",
        "PARTY_A_PHONE",
        "PARTY_B_ADDRESS",
        "PARTY_B_PHONE",
        "SIGN_YEAR",
        "SIGN_MONTH",
        "SIGN_DAY",
    )}
    mapping["{{PARTY_A_SIGNER}}"] = SIGNER_BLANK
    mapping["{{PARTY_B_SIGNER}}"] = SIGNER_BLANK
    _replace_paragraphs(document.paragraphs, mapping)
    for table in document.tables:
        _replace_in_table(table, mapping)
    for section in document.sections:
        _replace_paragraphs(section.header.paragraphs, mapping)
        _replace_paragraphs(section.footer.paragraphs, mapping)
        for table in section.header.tables:
            _replace_in_table(table, mapping)
        for table in section.footer.tables:
            _replace_in_table(table, mapping)
    _polish_filled_contract(document, key)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_placeholder_values(draft: dict) -> dict[str, str]:
    start_year, start_month, start_day = parse_date_parts(str(draft.get("startDate") or ""))
    end_year, end_month, end_day = parse_date_parts(str(draft.get("endDate") or ""))
    sign_year, sign_month, sign_day = parse_date_parts(str(draft.get("signDate") or draft.get("startDate") or ""))
    amount = draft.get("amount") or 0
    return {
        "PARTY_A_NAME": str(draft.get("partyAName") or "").strip(),
        "PARTY_B_NAME": str(draft.get("partyBName") or "").strip(),
        "SHOP_NAME": shop_display_name(str(draft.get("shopName") or "")),
        "SHOP_CATEGORY": format_shop_category(draft),
        "LISTING_LIMIT": str(draft.get("listingLimit") or "").strip(),
        "START_YEAR": start_year,
        "START_MONTH": start_month,
        "START_DAY": start_day,
        "END_YEAR": end_year,
        "END_MONTH": end_month,
        "END_DAY": end_day,
        "AMOUNT": format_amount(amount),
        "AMOUNT_CN": amount_to_chinese(amount),
        "ACCOUNT_NAME": str(draft.get("accountName") or "").strip(),
        "PARTY_A_SIGNER": SIGNER_BLANK,
        "PARTY_A_ADDRESS": str(draft.get("partyAAddress") or "").strip(),
        "PARTY_A_PHONE": str(draft.get("partyAPhone") or "").strip(),
        "PARTY_B_SIGNER": SIGNER_BLANK,
        "PARTY_B_ADDRESS": str(draft.get("partyBAddress") or "").strip(),
        "PARTY_B_PHONE": str(draft.get("partyBPhone") or "").strip(),
        "SIGN_YEAR": sign_year,
        "SIGN_MONTH": sign_month,
        "SIGN_DAY": sign_day,
    }


def missing_required_fields(draft: dict, platform: str | None = None) -> list[str]:
    normalized = dict(draft)
    normalized["shopCategory"] = format_shop_category(normalized)
    required = REQUIRED_FIELDS
    if not contract_needs_collection_account(platform):
        required = tuple(key for key in REQUIRED_FIELDS if key not in ACCOUNT_FIELDS)
    missing = [key for key in required if not str(normalized.get(key) or "").strip()]
    try:
        amount = Decimal(str(draft.get("amount") or ""))
    except Exception:
        amount = Decimal("0")
    if amount <= 0:
        missing.append("amount")
    start = parse_date_parts(str(draft.get("startDate") or ""))
    end = parse_date_parts(str(draft.get("endDate") or ""))
    sign = parse_date_parts(str(draft.get("signDate") or ""))
    if not all(start):
        missing.append("startDate")
    if not all(end):
        missing.append("endDate")
    if not all(sign):
        missing.append("signDate")
    return list(dict.fromkeys(missing))


DRAFT_OVERRIDE_KEYS = (
    "partyAName",
    "partyAAddress",
    "partyAPhone",
    "partyBName",
    "partyBAddress",
    "partyBPhone",
    "shopName",
    "shopCategory",
    "shopCategories",
    "listingLimit",
    "startDate",
    "endDate",
    "signDate",
    "amount",
    "accountName",
)


def apply_draft_overrides(base: dict, payload: dict | None) -> dict:
    merged = dict(base)
    if not isinstance(payload, dict):
        merged["shopCategory"] = format_shop_category(merged)
        return merged
    for key in DRAFT_OVERRIDE_KEYS:
        if key in payload:
            merged[key] = payload.get(key)
    merged["shopCategory"] = format_shop_category(merged)
    return merged


def build_default_draft(
    project: dict,
    customer: dict | None,
    own_company: dict | None,
    service_orders: list[dict],
    account_items: list[dict],
) -> dict:
    platform = normalize_contract_platform(str(project.get("sourcePlatform") or ""))
    defaults = PLATFORM_DEFAULTS.get(platform, {"shopCategories": ["服务类"], "listingLimit": "300", "amount": "3000"})
    customer = customer or {}
    own_company = own_company or {}
    party_a_name = str(customer.get("name") or "").strip()
    party_b_name = str(own_company.get("name") or "").strip()
    start_date = str(project.get("serviceStartDate") or "")[:10]
    end_date = str(project.get("serviceEndDate") or "")[:10]
    shop_categories = list(defaults["shopCategories"])
    return {
        "partyAName": party_a_name,
        "partyAAddress": str(customer.get("address") or "").strip(),
        "partyAPhone": str(customer.get("phone") or "").strip(),
        "partyBName": party_b_name,
        "partyBAddress": str(own_company.get("address") or "").strip(),
        "partyBPhone": str(own_company.get("phone") or "").strip(),
        "shopName": shop_name_from_company(party_a_name),
        "shopCategories": shop_categories,
        "shopCategory": format_shop_category({"shopCategories": shop_categories}),
        "listingLimit": defaults["listingLimit"],
        "startDate": start_date,
        "endDate": end_date,
        "signDate": start_date,
        "amount": _project_or_default_amount(project, defaults),
        "bankName": template_bank_account(platform)[0],
        "accountName": party_b_name,
        "accountNumber": template_bank_account(platform)[1],
    }


def find_collection_account(project: dict, service_orders: list[dict], account_items: list[dict]) -> dict | None:
    own_id = str(project.get("ownCompanyOrgId") or "")
    enabled_items = [
        item for item in account_items
        if isinstance(item, dict) and item.get("enabled") is not False and (
            item.get("accountNumber") or item.get("value")
        )
    ]
    orders = [
        order for order in service_orders
        if isinstance(order, dict)
        and order.get("projectId") == project.get("id")
        and order.get("status") not in {"取消", "退款"}
    ]
    orders.sort(key=lambda item: str(item.get("updatedAt") or item.get("createdAt") or ""), reverse=True)
    raw_account = next((str(order.get("collectionAccount") or "").strip() for order in orders if order.get("collectionAccount")), "")
    matched = _match_account_item(raw_account, enabled_items) if raw_account else None
    if matched:
        return matched
    own_accounts = [
        item for item in enabled_items
        if not own_id or str(item.get("ownCompanyOrgId") or "") == own_id
    ]
    if own_accounts:
        return own_accounts[0]
    return None


def _match_account_item(account: str, items: list[dict]) -> dict | None:
    by_id = next((item for item in items if item.get("id") == account), None)
    if by_id:
        return by_id
    numbered = [
        item for item in items
        if (item.get("accountNumber") and (account == item.get("accountNumber") or account.find(str(item.get("accountNumber"))) >= 0))
        or (item.get("value") and account == item.get("value"))
    ]
    if len(numbered) == 1:
        return numbered[0]
    if numbered:
        return numbered[0]
    return None


def contract_filename(draft: dict, platform: str, output_format: str = "word") -> str:
    year, _, _ = parse_date_parts(str(draft.get("signDate") or draft.get("startDate") or ""))
    if not year:
        from datetime import date
        year = str(date.today().year)
    company = str(draft.get("partyAName") or draft.get("shopName") or "服务合同").strip()
    safe = re.sub(r'[\\/:*?"<>|]+', "_", company) or "服务合同"
    ext = "pdf" if normalize_contract_format(output_format) == "pdf" else "docx"
    return f"{year}{platform}服务合同-{safe}.{ext}"


def render_contract_file(platform: str, draft: dict, output_format: str = "word") -> tuple[bytes, str, str]:
    fmt = normalize_contract_format(output_format)
    if fmt not in {"word", "pdf"}:
        raise ValueError("合同格式仅支持 Word 或 PDF。")
    content = fill_contract_document(platform, build_placeholder_values(draft))
    filename = contract_filename(draft, platform, fmt)
    if fmt == "pdf":
        return convert_docx_to_pdf(content), PDF_CONTENT_TYPE, filename
    return _refresh_docx_page_count(content), WORD_CONTENT_TYPE, filename


def _refresh_docx_page_count(docx_bytes: bytes) -> bytes:
    soffice = _find_soffice()
    if not soffice:
        return docx_bytes
    try:
        with tempfile.TemporaryDirectory(prefix="contract-pages-") as tmp:
            pdf = _run_soffice_pdf(soffice, _fit_contract_to_word_grid(docx_bytes), Path(tmp) / "count")
        total = _pdf_page_count(pdf)
    except Exception:
        logger.warning("Could not refresh contract page count", exc_info=True)
        return docx_bytes
    if total < 1 or _cached_numpages_value(docx_bytes) == total:
        return docx_bytes
    return apply_known_page_count(docx_bytes, total)


def convert_docx_to_pdf(docx_bytes: bytes) -> bytes:
    soffice = _find_soffice()
    if not soffice:
        raise ValueError("当前环境无法生成 PDF，请改选 Word 格式。")
    fitted = _fit_contract_to_word_grid(docx_bytes)
    with tempfile.TemporaryDirectory(prefix="contract-pdf-") as tmp:
        tmp_path = Path(tmp)
        pdf = _run_soffice_pdf(soffice, fitted, tmp_path / "pass1")
        total = _pdf_page_count(pdf)
        if total < 1 or _cached_numpages_value(docx_bytes) == total:
            return pdf
        patched = apply_known_page_count(fitted, total)
        return _run_soffice_pdf(soffice, patched, tmp_path / "pass2")


def _document_line_pitch_twips(document: Document) -> int:
    for section in document.sections:
        grid = section._sectPr.find(qn("w:docGrid"))
        if grid is None:
            continue
        raw = grid.get(qn("w:linePitch"))
        if raw and str(raw).isdigit():
            pitch = int(raw)
            if pitch > 0:
                return pitch
    return 312


def _paragraph_max_half_points(paragraph) -> int:
    sizes = [21]
    for node in paragraph.iter(qn("w:sz")):
        raw = node.get(qn("w:val"))
        if raw and str(raw).isdigit():
            sizes.append(int(raw))
    return max(sizes)


def _paragraph_has_drawing(paragraph) -> bool:
    return any(True for _ in paragraph.iter(qn("w:drawing"))) or any(True for _ in paragraph.iter(qn("w:pict")))


def _set_paragraph_exact_spacing(paragraph, pitch_twips: int) -> None:
    if _paragraph_has_drawing(paragraph):
        return
    font_twips = _paragraph_max_half_points(paragraph) * 10
    line_twips = max(pitch_twips, font_twips + 40)
    p_pr = paragraph.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = OxmlElement("w:pPr")
        paragraph.insert(0, p_pr)
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:line"), str(line_twips))
    spacing.set(qn("w:lineRule"), "exact")


def _keep_signature_block_together(document: Document) -> None:
    start = next(
        (paragraph for paragraph in document.element.body.iter(qn("w:p")) if "以下无正文" in "".join(paragraph.itertext())),
        None,
    )
    if start is None:
        return
    p_pr = start.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = OxmlElement("w:pPr")
        start.insert(0, p_pr)
    page_break = p_pr.find(qn("w:pageBreakBefore"))
    if page_break is None:
        page_break = OxmlElement("w:pageBreakBefore")
        p_pr.append(page_break)
    page_break.set(qn("w:val"), "true")


def _fit_contract_to_word_grid(docx_bytes: bytes) -> bytes:
    """LibreOffice ignores East-Asian docGrid, so 1.5 line spacing plus CJK fallback
    fonts overflow Word's 4-page layout. Lock body lines to the template grid and
    keep the signature block on the last page."""
    document = Document(BytesIO(docx_bytes))
    pitch = _document_line_pitch_twips(document)
    for paragraph in document.element.body.iter(qn("w:p")):
        _set_paragraph_exact_spacing(paragraph, pitch)
    _keep_signature_block_together(document)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _run_soffice_pdf(soffice: str, docx_bytes: bytes, outdir: Path) -> bytes:
    outdir.mkdir(parents=True, exist_ok=True)
    source = outdir / "contract.docx"
    source.write_bytes(docx_bytes)
    profile = outdir / "lo-profile"
    env = os.environ.copy()
    env["HOME"] = str(outdir)
    env["SAL_USE_VCLPLUGIN"] = "svp"
    env.setdefault("LC_PAPER", "a4")
    env.setdefault("PAPERSIZE", "a4")
    command = [
        soffice,
        "--headless",
        "--norestore",
        "--nolockcheck",
        "--nodefault",
        "--nofirststartwizard",
        f"-env:UserInstallation={profile.resolve().as_uri()}",
        "--convert-to",
        "pdf:writer_pdf_Export",
        "--outdir",
        str(outdir),
        str(source),
    ]
    try:
        result = subprocess.run(command, capture_output=True, text=True, timeout=90, check=False, env=env)
    except subprocess.TimeoutExpired as exc:
        logger.warning("LibreOffice PDF conversion timed out: %s", exc)
        raise ValueError("合同 PDF 生成超时，请稍后重试或改选 Word 格式。") from exc
    pdf = outdir / "contract.pdf"
    if not pdf.is_file():
        logger.warning("LibreOffice PDF conversion failed: code=%s stderr=%s stdout=%s", result.returncode, result.stderr, result.stdout)
        raise ValueError("合同 PDF 生成失败，请稍后重试或改选 Word 格式。")
    return pdf.read_bytes()


def _pdf_page_count(pdf: bytes) -> int:
    totals: list[int] = []
    for match in re.finditer(rb"/Type\s*/Pages\b(.{0,240})", pdf, re.S):
        count = re.search(rb"/Count\s+(\d+)", match.group(1))
        if count:
            totals.append(int(count.group(1)))
    if totals:
        return max(totals)
    return len(re.findall(rb"/Type\s*/Page(?!s)", pdf))


def _cached_numpages_value(docx_bytes: bytes) -> int | None:
    values: list[int] = []
    with ZipFile(BytesIO(docx_bytes)) as archive:
        for name in archive.namelist():
            if not (name.startswith("word/footer") and name.endswith(".xml")):
                continue
            root = etree.fromstring(archive.read(name))
            cached = _numpages_field_result(root)
            if cached is not None:
                values.append(cached)
    return values[0] if values else None


def apply_known_page_count(docx_bytes: bytes, total: int) -> bytes:
    if total < 1:
        return docx_bytes
    text = str(total)
    output = BytesIO()
    with ZipFile(BytesIO(docx_bytes)) as source, ZipFile(output, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename.startswith("word/footer") and info.filename.endswith(".xml"):
                root = etree.fromstring(payload)
                _set_numpages_results(root, text)
                payload = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif info.filename == "docProps/app.xml":
                payload = re.sub(br"<Pages>\d+</Pages>", f"<Pages>{total}</Pages>".encode("ascii"), payload, count=1)
            new_info = ZipInfo(filename=info.filename, date_time=info.date_time)
            new_info.compress_type = info.compress_type
            target.writestr(new_info, payload)
    return output.getvalue()


def _find_soffice() -> str | None:
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    for path in _SOFFICE_CANDIDATES:
        if path.is_file():
            return str(path)
    return None


def _project_or_default_amount(project: dict, defaults: dict) -> str:
    try:
        amount = Decimal(str(project.get("totalAmount") or 0))
    except Exception:
        amount = Decimal("0")
    if amount > 0:
        return format_amount(amount)
    return str(defaults.get("amount") or "3000")


def _integer_to_chinese(value: int) -> str:
    if value == 0:
        return "零"
    units = ["", "万", "亿"]
    parts: list[str] = []
    unit_index = 0
    pending_zero = False
    number = value
    while number > 0:
        number, section = divmod(number, 10000)
        section_text = _section_to_chinese(section)
        if section == 0:
            if parts:
                pending_zero = True
        else:
            if pending_zero:
                parts.append("零")
                pending_zero = False
            parts.append(section_text + units[unit_index])
        unit_index += 1
    return "".join(reversed(parts))


def _section_to_chinese(value: int) -> str:
    digits = [value // 1000, (value // 100) % 10, (value // 10) % 10, value % 10]
    names = ["仟", "佰", "拾", ""]
    result: list[str] = []
    seen_nonzero = False
    pending_zero = False
    for digit, name in zip(digits, names):
        if digit == 0:
            if seen_nonzero:
                pending_zero = True
            continue
        if pending_zero:
            result.append("零")
            pending_zero = False
        seen_nonzero = True
        result.append(CN_DIGITS[digit] + name)
    return "".join(result)


def _replace_paragraphs(paragraphs, mapping: dict[str, str]) -> None:
    for paragraph in paragraphs:
        original = paragraph.text
        if not any(token in original for token in mapping):
            continue
        parts: list[tuple[str, bool]] = []
        last = 0
        for match in TOKEN_RE.finditer(original):
            if match.start() > last:
                parts.append((original[last:match.start()], False))
            token = match.group(0)
            value = mapping.get(token, token)
            if token in SIGNER_TOKENS and not str(value).strip():
                value = SIGNER_BLANK
            parts.append((value, True))
            last = match.end()
        if last < len(original):
            parts.append((original[last:], False))
        _rewrite_paragraph_runs(paragraph, parts)


def _rewrite_paragraph_runs(paragraph, parts: list[tuple[str, bool]]) -> None:
    template_rpr = None
    if paragraph.runs:
        existing = paragraph.runs[0]._r.find(qn("w:rPr"))
        if existing is not None:
            template_rpr = deepcopy(existing)
    paragraph_el = paragraph._p
    for child in list(paragraph_el):
        if child.tag == qn("w:r"):
            paragraph_el.remove(child)
    for text, underlined in parts:
        run = paragraph.add_run(text)
        if template_rpr is not None:
            existing = run._r.find(qn("w:rPr"))
            if existing is not None:
                run._r.remove(existing)
            run._r.insert(0, deepcopy(template_rpr))
        if underlined:
            run.underline = True


def _replace_in_table(table, mapping: dict[str, str]) -> None:
    for row in table.rows:
        for cell in row.cells:
            _replace_paragraphs(cell.paragraphs, mapping)
            for nested in cell.tables:
                _replace_in_table(nested, mapping)


PAYMENT_QR_PATH = TEMPLATE_DIR / "payment_qr.png"


def _polish_filled_contract(document: Document, platform: str) -> None:
    # Word shows the brand slogan + page no. in one line via a narrow VML textbox;
    # LibreOffice PDF keeps the 2in box and wraps spaces, so rebuild with tab stops.
    _normalize_contract_footers(document)
    _ensure_document_page_fields(document)
    _enable_update_fields(document)
    if contract_needs_payment_qr(platform):
        _ensure_account_qr(document)
    else:
        _remove_floating_vml_images(document.element.body)


def _normalize_contract_footers(document: Document) -> None:
    for section in document.sections:
        usable_emu = int(section.page_width) - int(section.left_margin) - int(section.right_margin)
        # OOXML tab positions are twips; Length subtraction may already be an int (EMU).
        usable_twips = max(int(round(usable_emu / 635)), 1)
        _rewrite_footer_element(section.footer._element, usable_twips)
        if section.different_first_page_header_footer:
            _rewrite_footer_element(section.first_page_footer._element, usable_twips)
        if section._sectPr.find(qn("w:evenAndOddHeaders")) is not None:
            _rewrite_footer_element(section.even_page_footer._element, usable_twips)


def _rewrite_footer_element(footer_el, usable_twips: int) -> None:
    for child in list(footer_el):
        tag = child.tag
        if tag in {qn("w:p"), qn("w:tbl"), qn("w:sdt")}:
            footer_el.remove(child)
    footer_el.append(_build_tabbed_footer_paragraph(usable_twips))


def _build_tabbed_footer_paragraph(usable_twips: int):
    paragraph = OxmlElement("w:p")
    p_pr = OxmlElement("w:pPr")

    p_bdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "6")
    top.set(qn("w:space"), "8")
    top.set(qn("w:color"), "7F7F7F")
    p_bdr.append(top)
    p_pr.append(p_bdr)

    tabs = OxmlElement("w:tabs")
    for value, pos in (("center", usable_twips // 2), ("right", usable_twips)):
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), value)
        tab.set(qn("w:pos"), str(pos))
        tabs.append(tab)
    p_pr.append(tabs)
    paragraph.append(p_pr)

    _append_footer_text_run(paragraph, FOOTER_LEFT_TEXT)
    _append_footer_tab_run(paragraph)
    _append_footer_page_field(paragraph, "PAGE", "1")
    _append_footer_text_run(paragraph, " / ", preserve=True)
    _append_footer_page_field(paragraph, "NUMPAGES", "4")
    _append_footer_tab_run(paragraph)
    _append_footer_text_run(paragraph, FOOTER_RIGHT_TEXT)
    return paragraph


def _append_footer_run_props(run) -> None:
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:hint"), "eastAsia")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), "18")
    r_pr.append(fonts)
    r_pr.append(sz)
    r_pr.append(sz_cs)
    run.insert(0, r_pr)


def _append_footer_text_run(paragraph, text: str, *, preserve: bool = False) -> None:
    run = OxmlElement("w:r")
    _append_footer_run_props(run)
    node = OxmlElement("w:t")
    if preserve or text[:1].isspace() or text[-1:].isspace():
        node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    node.text = text
    run.append(node)
    paragraph.append(run)


def _append_footer_tab_run(paragraph) -> None:
    run = OxmlElement("w:r")
    _append_footer_run_props(run)
    run.append(OxmlElement("w:tab"))
    paragraph.append(run)


def _append_footer_page_field(paragraph, instr: str, placeholder: str) -> None:
    begin_run = OxmlElement("w:r")
    _append_footer_run_props(begin_run)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    begin_run.append(begin)
    paragraph.append(begin_run)

    instr_run = OxmlElement("w:r")
    _append_footer_run_props(instr_run)
    instr_el = OxmlElement("w:instrText")
    instr_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr_el.text = f" {instr} "
    instr_run.append(instr_el)
    paragraph.append(instr_run)

    sep_run = OxmlElement("w:r")
    _append_footer_run_props(sep_run)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    sep_run.append(separate)
    paragraph.append(sep_run)

    text_run = OxmlElement("w:r")
    _append_footer_run_props(text_run)
    text = OxmlElement("w:t")
    text.text = placeholder
    text_run.append(text)
    paragraph.append(text_run)

    end_run = OxmlElement("w:r")
    _append_footer_run_props(end_run)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)
    paragraph.append(end_run)


def _enable_update_fields(document: Document) -> None:
    settings = document.settings.element
    if settings.find(qn("w:updateFields")) is not None:
        return
    node = OxmlElement("w:updateFields")
    node.set(qn("w:val"), "true")
    settings.append(node)


def _ensure_document_page_fields(document: Document) -> None:
    for section in document.sections:
        _ensure_footer_page_total_field(section.footer._element)
        if section.different_first_page_header_footer:
            _ensure_footer_page_total_field(section.first_page_footer._element)
        if section._sectPr.find(qn("w:evenAndOddHeaders")) is not None:
            _ensure_footer_page_total_field(section.even_page_footer._element)


def _ensure_footer_page_total_field(footer_el) -> None:
    if any((node.text or "").strip() == "NUMPAGES" for node in footer_el.iter(qn("w:instrText"))):
        _mark_page_fields_dirty(footer_el)
        return
    for paragraph in footer_el.iter(qn("w:p")):
        if _replace_static_total_with_numpages(paragraph):
            return


def _replace_static_total_with_numpages(paragraph) -> bool:
    runs = [child for child in paragraph if child.tag == qn("w:r")]
    index = 0
    while index < len(runs):
        if _fld_type(runs[index]) != "begin":
            index += 1
            continue
        end = _field_end_index(runs, index)
        instr = _field_instr(runs, index, end)
        if instr != "PAGE" or end is None:
            index = end + 1 if end is not None else index + 1
            continue
        next_index = end + 1
        if next_index >= len(runs):
            return False
        separator = _run_plain_text(runs[next_index])
        if separator.strip() == "/" and next_index + 1 < len(runs):
            total_run = runs[next_index + 1]
            total_text = _run_plain_text(total_run).strip()
            if total_text.isdigit():
                _replace_run_with_field(total_run, "NUMPAGES", total_text)
                return True
        match = re.fullmatch(r"(\s*/\s*)(\d+)(\s*)", separator)
        if match:
            _replace_run_with_field(runs[next_index], "NUMPAGES", match.group(2))
            return True
        index = next_index
    return False


def _replace_run_with_field(run, instr: str, placeholder: str) -> None:
    parent = run.getparent()
    if parent is None:
        return
    rpr = run.find(qn("w:rPr"))

    def make_run(*children):
        element = OxmlElement("w:r")
        if rpr is not None:
            element.append(deepcopy(rpr))
        for child in children:
            element.append(child)
        return element

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    replacements = [
        make_run(begin),
        make_run(instr_el),
        make_run(separate),
        make_run(text),
        make_run(end),
    ]
    insert_at = list(parent).index(run)
    parent.remove(run)
    for offset, item in enumerate(replacements):
        parent.insert(insert_at + offset, item)


def _set_numpages_results(root, value: str) -> None:
    for paragraph in root.iter(qn("w:p")):
        runs = [child for child in paragraph if child.tag == qn("w:r")]
        index = 0
        while index < len(runs):
            if _fld_type(runs[index]) != "begin":
                index += 1
                continue
            end = _field_end_index(runs, index)
            if end is None:
                index += 1
                continue
            if _field_instr(runs, index, end) == "NUMPAGES":
                _mark_fld_begin_dirty(runs[index])
                in_result = False
                for run in runs[index:end]:
                    if _fld_type(run) == "separate":
                        in_result = True
                        continue
                    if in_result:
                        texts = list(run.iter(qn("w:t")))
                        if texts:
                            texts[0].text = value
                            for extra in texts[1:]:
                                extra.text = ""
            index = end + 1


def _numpages_field_result(root) -> int | None:
    for paragraph in root.iter(qn("w:p")):
        runs = [child for child in paragraph if child.tag == qn("w:r")]
        index = 0
        while index < len(runs):
            if _fld_type(runs[index]) != "begin":
                index += 1
                continue
            end = _field_end_index(runs, index)
            if end is None:
                index += 1
                continue
            if _field_instr(runs, index, end) == "NUMPAGES":
                in_result = False
                collected: list[str] = []
                for run in runs[index:end]:
                    if _fld_type(run) == "separate":
                        in_result = True
                        continue
                    if in_result:
                        collected.append(_run_plain_text(run))
                text = "".join(collected).strip()
                if text.isdigit():
                    return int(text)
            index = end + 1
    return None


def _mark_page_fields_dirty(root) -> None:
    for paragraph in root.iter(qn("w:p")):
        runs = [child for child in paragraph if child.tag == qn("w:r")]
        index = 0
        while index < len(runs):
            if _fld_type(runs[index]) != "begin":
                index += 1
                continue
            end = _field_end_index(runs, index)
            if end is not None and _field_instr(runs, index, end) in {"PAGE", "NUMPAGES"}:
                _mark_fld_begin_dirty(runs[index])
            index = end + 1 if end is not None else index + 1


def _mark_fld_begin_dirty(run) -> None:
    fld = run.find(qn("w:fldChar"))
    if fld is not None:
        fld.set(qn("w:dirty"), "true")


def _fld_type(run) -> str | None:
    fld = run.find(qn("w:fldChar"))
    if fld is None:
        return None
    return fld.get(qn("w:fldCharType"))


def _field_end_index(runs: list, start: int) -> int | None:
    depth = 0
    for index in range(start, len(runs)):
        kind = _fld_type(runs[index])
        if kind == "begin":
            depth += 1
        elif kind == "end":
            depth -= 1
            if depth == 0:
                return index
    return None


def _field_instr(runs: list, start: int, end: int | None) -> str:
    if end is None:
        return ""
    parts: list[str] = []
    for run in runs[start:end + 1]:
        instr = run.find(qn("w:instrText"))
        if instr is not None:
            parts.append(instr.text or "")
    return "".join(parts).strip()


def _run_plain_text(run) -> str:
    return "".join(node.text or "" for node in run.findall(qn("w:t")))


def _ensure_account_qr(document: Document) -> None:
    _remove_floating_vml_images(document.element.body)
    if not PAYMENT_QR_PATH.is_file():
        return
    table = next((item for item in document.tables if any("开户行" in cell.text for row in item.rows for cell in row.cells)), None)
    if table is None:
        return
    if len(table.rows[0].cells) < 2:
        _append_table_cell(table, "2400")
    cell = table.rows[0].cells[-1]
    for paragraph in cell.paragraphs:
        if paragraph.text.strip():
            continue
        for run in list(paragraph.runs):
            run._r.getparent().remove(run._r)
    target = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    target.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if not target.runs:
        target.add_run()
    target.runs[0].add_picture(str(PAYMENT_QR_PATH), width=Pt(63))
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _append_table_cell(table, width_dxa: str) -> None:
    row_el = table.rows[0]._tr
    tc = OxmlElement("w:tc")
    tc_pr = OxmlElement("w:tcPr")
    tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:w"), width_dxa)
    tc_w.set(qn("w:type"), "dxa")
    tc_pr.append(tc_w)
    tc.append(tc_pr)
    tc.append(OxmlElement("w:p"))
    row_el.append(tc)
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), width_dxa)
        grid.append(col)


def _remove_floating_vml_images(root) -> None:
    for pict in list(root.iter(qn("w:pict"))):
        if pict.find(".//{urn:schemas-microsoft-com:vml}imagedata") is None:
            continue
        parent = pict.getparent()
        if parent is not None:
            parent.remove(pict)
