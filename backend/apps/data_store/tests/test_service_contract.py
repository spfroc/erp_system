from io import BytesIO
from zipfile import ZipFile

from unittest.mock import patch
from urllib.parse import quote

from django.test import Client, SimpleTestCase, TestCase
from docx import Document

from apps.data_store.service_contract import (
    SIGNER_BLANK,
    amount_to_chinese,
    apply_known_page_count,
    build_default_draft,
    contract_filename,
    convert_docx_to_pdf,
    fill_contract_document,
    format_shop_category,
    has_service_contract_template,
    missing_required_fields,
    normalize_contract_format,
    render_contract_file,
    shop_name_from_company,
    _cached_numpages_value,
    _fit_contract_to_word_grid,
    _pdf_page_count,
)
from apps.data_store.views import save_app_data
from docx.oxml.ns import qn

SAMPLE_CONTRACT_VALUES = {
    "PARTY_A_NAME": "示例智能科技有限公司",
    "PARTY_B_NAME": "济南云集企业管理有限公司",
    "SHOP_NAME": "示例智能科技",
    "SHOP_CATEGORY": "货物类",
    "LISTING_LIMIT": "不限数量",
    "START_YEAR": "2026",
    "START_MONTH": "6",
    "START_DAY": "24",
    "END_YEAR": "2028",
    "END_MONTH": "6",
    "END_DAY": "23",
    "AMOUNT": "4000",
    "AMOUNT_CN": "肆仟元整",
    "BANK_NAME": "中信银行济南市中支行",
    "ACCOUNT_NAME": "济南云集企业管理有限公司",
    "ACCOUNT_NO": "8112501012401497344",
    "PARTY_A_SIGNER": "赵入驻",
    "PARTY_A_ADDRESS": "济南市高新区示例路 1 号",
    "PARTY_A_PHONE": "13065056229",
    "PARTY_B_SIGNER": "法人",
    "PARTY_B_ADDRESS": "济南市市中区示例路 8 号",
    "PARTY_B_PHONE": "0531-86099058",
    "SIGN_YEAR": "2026",
    "SIGN_MONTH": "6",
    "SIGN_DAY": "24",
}


def _document_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


class ServiceContractHelperTests(SimpleTestCase):
    def test_shop_name_strips_company_suffix(self):
        self.assertEqual(shop_name_from_company("山东壹知产数字科技有限公司"), "山东壹知产数字科技店")
        self.assertEqual(shop_name_from_company("济南示例有限责任公司"), "济南示例店")
        self.assertEqual(shop_name_from_company("示例智能科技店"), "示例智能科技店")

    def test_amount_to_chinese(self):
        self.assertEqual(amount_to_chinese(3000), "叁仟元整")
        self.assertEqual(amount_to_chinese("4000"), "肆仟元整")
        self.assertEqual(amount_to_chinese(108.3), "壹佰零捌元叁角")
        self.assertEqual(amount_to_chinese(10), "壹拾元整")

    def test_template_platforms(self):
        self.assertTrue(has_service_contract_template("泉E采"))
        self.assertTrue(has_service_contract_template("泉e采"))
        self.assertTrue(has_service_contract_template("青慧采"))
        self.assertFalse(has_service_contract_template("齐鲁云采"))

    def test_fill_qec_template(self):
        content = fill_contract_document("泉E采", {
            **SAMPLE_CONTRACT_VALUES,
            "SHOP_CATEGORY": "服务类",
            "LISTING_LIMIT": "300",
            "AMOUNT": "2000",
            "AMOUNT_CN": "贰仟元整",
        })
        self.assertGreater(len(content), 20000)
        self.assertTrue(content.startswith(b"PK"))
        document = Document(BytesIO(content))
        runs = [run for paragraph in document.paragraphs for run in paragraph.runs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        runs.extend(paragraph.runs)
        self.assertTrue(any("示例智能科技" in (run.text or "") for run in runs))
        self.assertTrue(any(run.underline and "示例智能科技" in (run.text or "") for run in runs))
        footer_xml = "".join(section.footer._element.xml for section in document.sections)
        self.assertIn("链接平台", footer_xml)
        self.assertIn("链接厂商", footer_xml)
        self.assertIn("链接数据", footer_xml)
        self.assertIn("创造价值", footer_xml)
        self.assertIn("NUMPAGES", footer_xml)
        self.assertIn("w:tab", footer_xml)
        self.assertNotIn("v:textbox", footer_xml)
        self.assertNotIn("文本框", footer_xml)
        self.assertEqual(len(document.sections[0].footer.tables), 0)
        account_table = next(table for table in document.tables if any("开户行" in cell.text for row in table.rows for cell in row.cells))
        self.assertGreaterEqual(len(account_table.rows[0].cells), 2)
        signer_runs = [run for run in runs if run.text == SIGNER_BLANK]
        self.assertTrue(signer_runs)
        self.assertTrue(all(run.underline for run in signer_runs))
        self.assertIn("乙方收款账户", _document_text(document))
        self.assertTrue(list(document.element.body.iter(qn("a:blip"))))

    def test_fill_qhc_template_keeps_account_without_qr(self):
        content = fill_contract_document("青慧采", SAMPLE_CONTRACT_VALUES)
        document = Document(BytesIO(content))
        text = _document_text(document)
        self.assertIn("青慧采", text)
        self.assertNotIn("泉e采", text)
        self.assertIn("乙方收款账户", text)
        self.assertIn("开户行", text)
        self.assertIn("广发银行股份有限公司济南分行营业部", text)
        self.assertIn("9550880219307000169", text)
        self.assertFalse(list(document.element.body.iter(qn("a:blip"))))
        self.assertFalse(list(document.element.body.iter("{urn:schemas-microsoft-com:vml}imagedata")))
        account_table = next(table for table in document.tables if any("开户行" in cell.text for row in table.rows for cell in row.cells))
        self.assertEqual(len(account_table.rows[0].cells), 1)

    def test_fill_keeps_template_bank_and_account(self):
        values = {
            **SAMPLE_CONTRACT_VALUES,
            "BANK_NAME": "错误开户行",
            "ACCOUNT_NO": "000000",
        }
        expected = {
            "泉E采": ("中信银行济南市中支行", "8112501012401497344"),
            "青慧采": ("广发银行股份有限公司济南分行营业部", "9550880219307000169"),
        }
        for platform, (bank, account) in expected.items():
            content = fill_contract_document(platform, values)
            text = _document_text(Document(BytesIO(content)))
            self.assertIn(bank, text)
            self.assertIn(account, text)
            self.assertNotIn("错误开户行", text)
            self.assertNotIn("000000", text)

    def test_default_draft_and_missing_fields(self):
        draft = build_default_draft(
            {
                "id": "pj-1",
                "sourcePlatform": "泉E采",
                "ownCompanyOrgId": "org-own",
                "serviceStartDate": "2026-06-24",
                "serviceEndDate": "2027-06-23",
                "totalAmount": 3000,
            },
            {"name": "示例智能科技有限公司", "address": "客户地址", "phone": "13000000000", "legalPerson": "甲方法人"},
            {"name": "济南云集企业管理有限公司", "address": "乙方地址", "phone": "0531-86099058", "legalPerson": "乙方法人"},
            [{"id": "svc-1", "projectId": "pj-1", "status": "已付款", "collectionAccount": "acc-1", "createdAt": "2026-06-24"}],
            [{
                "id": "acc-1",
                "enabled": True,
                "ownCompanyOrgId": "org-own",
                "bankName": "中信银行济南市中支行",
                "accountNumber": "8112501012401497344",
            }],
        )
        self.assertEqual(draft["shopName"], "示例智能科技店")
        self.assertEqual(draft["shopCategories"], ["服务类"])
        self.assertEqual(draft["shopCategory"], "（政企专用）服务/工程类")
        self.assertEqual(draft["listingLimit"], "300")
        self.assertEqual(draft["amount"], "3000")
        self.assertEqual(draft["signDate"], "2026-06-24")
        self.assertEqual(missing_required_fields(draft), [])
        self.assertEqual(
            format_shop_category({"shopCategories": ["服务类", "产品类"]}),
            "（政企专用）服务/工程类/货物类",
        )
        qhc = build_default_draft(
            {**{
                "id": "pj-1",
                "sourcePlatform": "青慧采",
                "ownCompanyOrgId": "org-own",
                "serviceStartDate": "2026-06-24",
                "serviceEndDate": "2027-06-23",
                "totalAmount": 2000,
            }},
            {"name": "示例智能科技有限公司", "address": "客户地址", "phone": "13000000000"},
            {"name": "济南云集企业管理有限公司", "address": "乙方地址", "phone": "0531-86099058"},
            [{"id": "svc-1", "projectId": "pj-1", "status": "已付款", "collectionAccount": "acc-1", "createdAt": "2026-06-24"}],
            [{
                "id": "acc-1",
                "enabled": True,
                "ownCompanyOrgId": "org-own",
                "bankName": "中信银行济南市中支行",
                "accountNumber": "8112501012401497344",
            }],
        )
        self.assertEqual(qhc["amount"], "2000")
        self.assertEqual(qhc["shopCategory"], "货物类")
        qhc_default = build_default_draft(
            {
                "id": "pj-1",
                "sourcePlatform": "青慧采",
                "ownCompanyOrgId": "org-own",
                "serviceStartDate": "2026-06-24",
                "serviceEndDate": "2027-06-23",
                "totalAmount": 0,
            },
            {"name": "示例智能科技有限公司", "address": "客户地址", "phone": "13000000000"},
            {"name": "济南云集企业管理有限公司", "address": "乙方地址", "phone": "0531-86099058"},
            [{"id": "svc-1", "projectId": "pj-1", "status": "已付款", "collectionAccount": "acc-1", "createdAt": "2026-06-24"}],
            [{
                "id": "acc-1",
                "enabled": True,
                "ownCompanyOrgId": "org-own",
                "bankName": "中信银行济南市中支行",
                "accountNumber": "8112501012401497344",
            }],
        )
        self.assertEqual(qhc_default["amount"], "4000")
        self.assertEqual(qhc_default["bankName"], "广发银行股份有限公司济南分行营业部")
        self.assertEqual(qhc_default["accountNumber"], "9550880219307000169")
        missing = missing_required_fields({**qhc_default, "bankName": "", "accountName": "", "accountNumber": ""}, "青慧采")
        self.assertIn("accountName", missing)
        self.assertNotIn("bankName", missing)
        self.assertNotIn("accountNumber", missing)

    def test_contract_filename_uses_year_platform_and_party_a(self):
        self.assertEqual(
            contract_filename({"partyAName": "青岛康福润源贸易有限公司", "signDate": "2026-08-27"}, "青慧采"),
            "2026青慧采服务合同-青岛康福润源贸易有限公司.docx",
        )
        self.assertEqual(
            contract_filename({"partyAName": "青岛康福润源贸易有限公司", "signDate": "2026-08-27"}, "青慧采", "pdf"),
            "2026青慧采服务合同-青岛康福润源贸易有限公司.pdf",
        )
        self.assertEqual(normalize_contract_format("PDF"), "pdf")
        self.assertEqual(normalize_contract_format("docx"), "word")
        self.assertEqual(normalize_contract_format("xlsx"), "")

    def test_fill_keeps_boss_footer_and_uses_numpages(self):
        for platform in ("泉E采", "青慧采"):
            content = fill_contract_document(platform, SAMPLE_CONTRACT_VALUES)
            document = Document(BytesIO(content))
            footer_xml = document.sections[0].footer._element.xml
            self.assertIn("NUMPAGES", footer_xml)
            self.assertIn("PAGE", footer_xml)
            self.assertIn("w:tab", footer_xml)
            self.assertNotIn("v:textbox", footer_xml)
            self.assertEqual(_cached_numpages_value(content), 4)
            self.assertEqual(len(document.sections[0].footer.tables), 0)

    def test_apply_known_page_count_updates_footer_and_stats(self):
        content = fill_contract_document("青慧采", SAMPLE_CONTRACT_VALUES)
        patched = apply_known_page_count(content, 5)
        self.assertEqual(_cached_numpages_value(patched), 5)
        app_xml = ZipFile(BytesIO(patched)).read("docProps/app.xml").decode("utf-8")
        self.assertIn("<Pages>5</Pages>", app_xml)

    def test_fit_contract_locks_body_lines_to_document_grid(self):
        content = fill_contract_document("泉E采", SAMPLE_CONTRACT_VALUES)
        fitted = _fit_contract_to_word_grid(content)
        document = Document(BytesIO(fitted))
        rules = []
        lines = []
        for paragraph in document.element.body.iter(qn("w:p")):
            if any(True for _ in paragraph.iter(qn("w:drawing"))):
                continue
            spacing = None
            p_pr = paragraph.find(qn("w:pPr"))
            if p_pr is not None:
                spacing = p_pr.find(qn("w:spacing"))
            if spacing is None:
                continue
            rules.append(spacing.get(qn("w:lineRule")))
            lines.append(int(spacing.get(qn("w:line"))))
        self.assertTrue(rules)
        self.assertTrue(all(rule == "exact" for rule in rules))
        self.assertTrue(all(line >= 312 for line in lines))
        start = next(paragraph for paragraph in document.element.body.iter(qn("w:p")) if "以下无正文" in "".join(paragraph.itertext()))
        p_pr = start.find(qn("w:pPr"))
        page_break = p_pr.find(qn("w:pageBreakBefore")) if p_pr is not None else None
        self.assertIsNotNone(page_break)
        self.assertEqual(page_break.get(qn("w:val")), "true")

    def test_pdf_page_count_reads_pages_tree(self):
        pdf = (
            b"%PDF-1.4\n"
            b"1 0 obj<</Type /Catalog /Pages 2 0 R>>endobj\n"
            b"2 0 obj<</Type /Pages /Kids [3 0 R] /Count 5>>endobj\n"
            b"trailer<</Root 1 0 R>>\n%%EOF\n"
        )
        self.assertEqual(_pdf_page_count(pdf), 5)

    def test_convert_docx_to_pdf_refreshes_stale_total(self):
        content = fill_contract_document("青慧采", SAMPLE_CONTRACT_VALUES)
        pdf5 = (
            b"%PDF-1.4\n"
            b"1 0 obj<</Type /Catalog /Pages 2 0 R>>endobj\n"
            b"2 0 obj<</Type /Pages /Count 5 /Kids []>>endobj\n"
            b"trailer<</Root 1 0 R>>\n%%EOF\n"
        )
        calls: list[bytes] = []

        def fake_run(_soffice, data, _outdir):
            calls.append(data)
            return pdf5

        with patch("apps.data_store.service_contract._find_soffice", return_value="soffice"), patch(
            "apps.data_store.service_contract._run_soffice_pdf",
            side_effect=fake_run,
        ):
            result = convert_docx_to_pdf(content)
        self.assertEqual(result, pdf5)
        self.assertEqual(len(calls), 2)
        self.assertEqual(_cached_numpages_value(calls[1]), 5)

    def test_refresh_docx_page_count_updates_word_cache(self):
        pdf5 = (
            b"%PDF-1.4\n"
            b"1 0 obj<</Type /Catalog /Pages 2 0 R>>endobj\n"
            b"2 0 obj<</Type /Pages /Count 5 /Kids []>>endobj\n"
            b"trailer<</Root 1 0 R>>\n%%EOF\n"
        )
        with patch("apps.data_store.service_contract._find_soffice", return_value="soffice"), patch(
            "apps.data_store.service_contract._run_soffice_pdf",
            return_value=pdf5,
        ):
            refreshed, content_type, filename = render_contract_file("青慧采", {
                "partyAName": "示例智能科技有限公司",
                "partyBName": "济南云集企业管理有限公司",
                "shopName": "示例智能科技店",
                "shopCategory": "货物类",
                "listingLimit": "不限数量",
                "startDate": "2026-06-24",
                "endDate": "2028-06-23",
                "signDate": "2026-06-24",
                "amount": "4000",
                "bankName": "中信银行济南市中支行",
                "accountName": "济南云集企业管理有限公司",
                "accountNumber": "8112501012401497344",
                "partyAAddress": "济南",
                "partyAPhone": "130",
                "partyBAddress": "济南",
                "partyBPhone": "0531",
            }, "word")
        self.assertIn("wordprocessingml", content_type)
        self.assertTrue(filename.endswith(".docx"))
        self.assertEqual(_cached_numpages_value(refreshed), 5)

    def test_convert_docx_to_pdf_requires_soffice(self):
        with patch("apps.data_store.service_contract._find_soffice", return_value=None):
            with self.assertRaises(ValueError):
                convert_docx_to_pdf(b"PK fake")


class ServiceContractApiTests(TestCase):
    def setUp(self):
        self.client = Client()
        save_app_data({
            "users": [{
                "id": "user-admin",
                "name": "管理员",
                "phone": "13800000000",
                "password": "123456",
                "status": "启用",
                "roleIds": ["role-admin"],
            }],
            "roles": [{
                "id": "role-admin",
                "name": "管理员",
                "permissions": ["project.view", "project.edit"],
            }],
            "organizations": [
                {
                    "id": "org-a",
                    "name": "示例智能科技有限公司",
                    "entityType": "企业",
                    "roles": ["平台入驻商"],
                    "status": "正常",
                    "address": "济南市高新区示例路 1 号",
                    "phone": "13065056229",
                    "legalPerson": "赵入驻",
                    "createdAt": "2026-08-01T10:00:00",
                },
                {
                    "id": "org-own",
                    "name": "济南云集企业管理有限公司",
                    "entityType": "企业",
                    "roles": ["自有主体"],
                    "status": "正常",
                    "address": "山东省济南市市中区大观园街道纬四路4号3楼346室",
                    "phone": "0531-86099058",
                    "legalPerson": "乙方法人",
                    "createdAt": "2026-08-01T10:00:00",
                },
            ],
            "projects": [{
                "id": "proj-service",
                "projectNo": "FW-001",
                "name": "泉E采服务项目",
                "type": "服务项目",
                "customerOrgId": "org-a",
                "ownCompanyOrgId": "org-own",
                "sourcePlatform": "泉E采",
                "owner": "管理员",
                "status": "服务中",
                "contractStatus": "待签",
                "paymentStatus": "待付款",
                "purchaseStatus": "无需采购",
                "invoiceStatus": "待开",
                "totalAmount": 3000,
                "totalCost": 0,
                "serviceStartDate": "2026-06-24",
                "serviceEndDate": "2027-06-23",
                "updatedAt": "2026-08-05T10:00:00",
                "createdAt": "2026-08-05T10:00:00",
            }],
            "serviceOrders": [{
                "id": "svc-1",
                "orderNo": "SV-001",
                "projectId": "proj-service",
                "serviceId": "prd-1",
                "customerOrgId": "org-a",
                "amount": 3000,
                "status": "已付款",
                "collectionAccount": "acc-1",
                "createdAt": "2026-08-05T10:00:00",
                "updatedAt": "2026-08-05T10:00:00",
            }],
            "optionGroups": [{
                "id": "opt-collection",
                "key": "collection_account",
                "name": "收款账号",
                "usages": [],
                "items": [{
                    "id": "acc-1",
                    "label": "中信银行",
                    "value": "8112501012401497344",
                    "enabled": True,
                    "sortOrder": 1,
                    "ownCompanyOrgId": "org-own",
                    "bankName": "中信银行济南市中支行",
                    "accountNumber": "8112501012401497344",
                }],
            }],
        })
        login = self.client.post(
            "/api/auth/login/",
            data={"identifier": "13800000000", "password": "123456"},
            content_type="application/json",
        )
        self.assertEqual(login.status_code, 200)

    def test_generate_contract_downloads_docx(self):
        response = self.client.post(
            "/api/projects/proj-service/generate-contract/",
            data={},
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            response["Content-Type"],
        )
        self.assertTrue(response.content.startswith(b"PK"))
        self.assertIn(quote("2026泉E采服务合同-示例智能科技有限公司.docx"), response["Content-Disposition"])

    def test_generate_standalone_contract_downloads_docx(self):
        response = self.client.post(
            "/api/contracts/generate/",
            data={
                "platform": "青慧采",
                "partyAName": "青岛康福润源贸易有限公司",
                "partyAAddress": "青岛市示例路 1 号",
                "partyAPhone": "13000000000",
                "partyBName": "济南云集企业管理有限公司",
                "partyBAddress": "济南市市中区示例路 8 号",
                "partyBPhone": "0531-86099058",
                "shopName": "青岛康福润源贸易",
                "shopCategories": ["产品类"],
                "listingLimit": "不限数量",
                "startDate": "2026-08-27",
                "endDate": "2027-08-31",
                "signDate": "2026-08-27",
                "amount": "4000",
                "bankName": "广发银行股份有限公司济南分行营业部",
                "accountName": "济南云集企业管理有限公司",
                "accountNumber": "9550880219307000169",
            },
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.content.startswith(b"PK"))
        self.assertIn(quote("2026青慧采服务合同-青岛康福润源贸易有限公司.docx"), response["Content-Disposition"])
        document = Document(BytesIO(response.content))
        text = _document_text(document)
        self.assertIn("青慧采", text)
        self.assertIn("乙方收款账户", text)
        self.assertIn("开户行", text)
        self.assertIn("广发银行股份有限公司济南分行营业部", text)
        self.assertIn("9550880219307000169", text)

    @patch("apps.data_store.service_contract.convert_docx_to_pdf", return_value=b"%PDF-1.4 mock")
    def test_generate_standalone_contract_downloads_pdf(self, _convert):
        response = self.client.post(
            "/api/contracts/generate/",
            data={
                "platform": "青慧采",
                "format": "pdf",
                "partyAName": "青岛康福润源贸易有限公司",
                "partyAAddress": "青岛市示例路 1 号",
                "partyAPhone": "13000000000",
                "partyBName": "济南云集企业管理有限公司",
                "partyBAddress": "济南市市中区示例路 8 号",
                "partyBPhone": "0531-86099058",
                "shopName": "青岛康福润源贸易",
                "shopCategories": ["产品类"],
                "listingLimit": "不限数量",
                "startDate": "2026-08-27",
                "endDate": "2027-08-31",
                "signDate": "2026-08-27",
                "amount": "4000",
                "bankName": "中信银行济南市中支行",
                "accountName": "济南云集企业管理有限公司",
                "accountNumber": "8112501012401497344",
            },
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("application/pdf", response["Content-Type"])
        self.assertTrue(response.content.startswith(b"%PDF"))
        self.assertIn(quote("2026青慧采服务合同-青岛康福润源贸易有限公司.pdf"), response["Content-Disposition"])
        _convert.assert_called_once()

    def test_generate_contract_rejects_unknown_format(self):
        response = self.client.post(
            "/api/contracts/generate/",
            data={
                "platform": "青慧采",
                "format": "xlsx",
                "partyAName": "青岛康福润源贸易有限公司",
                "partyAAddress": "青岛市示例路 1 号",
                "partyAPhone": "13000000000",
                "partyBName": "济南云集企业管理有限公司",
                "partyBAddress": "济南市市中区示例路 8 号",
                "partyBPhone": "0531-86099058",
                "shopName": "青岛康福润源贸易",
                "shopCategories": ["产品类"],
                "listingLimit": "不限数量",
                "startDate": "2026-08-27",
                "endDate": "2027-08-31",
                "signDate": "2026-08-27",
                "amount": "4000",
                "bankName": "中信银行济南市中支行",
                "accountName": "济南云集企业管理有限公司",
                "accountNumber": "8112501012401497344",
            },
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("Word 或 PDF", response.json()["detail"])

    def test_generate_contract_rejects_platform_without_template(self):
        save_app_data({
            "projects": [{
                "id": "proj-service",
                "projectNo": "FW-001",
                "name": "齐鲁云采服务项目",
                "type": "服务项目",
                "customerOrgId": "org-a",
                "ownCompanyOrgId": "org-own",
                "sourcePlatform": "齐鲁云采",
                "owner": "管理员",
                "status": "服务中",
                "contractStatus": "待签",
                "paymentStatus": "待付款",
                "purchaseStatus": "无需采购",
                "invoiceStatus": "待开",
                "totalAmount": 3000,
                "totalCost": 0,
                "serviceStartDate": "2026-06-24",
                "serviceEndDate": "2027-06-23",
                "updatedAt": "2026-08-05T10:00:00",
                "createdAt": "2026-08-05T10:00:00",
            }],
        }, partial=True)
        response = self.client.post(
            "/api/projects/proj-service/generate-contract/",
            data={},
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("合同模板", response.json()["detail"])
