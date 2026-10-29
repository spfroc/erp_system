from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Optional
from urllib.parse import urlparse

from sqlalchemy import case, select
from tqdm import tqdm

from clients.http_session import proxy_url
from clients.jd_quan import JdQuanClient, is_jd_site
from config import settings
from db import NoticeAttachment, NoticeGoods, NoticeParsed, NoticeRaw, dumps, get_session
from pipeline.extractors.jd_attachment import (
    goods_are_placeholder,
    should_replace_goods_with_attachment,
)
from pipeline.parse_llm import (
    enrich_from_raw,
    extract_attachment_with_local_model,
    parsed_row_to_dict,
    save_parsed,
)
from utils import ensure_https
from utils import now as now_shanghai


def _ext_from_name(name: str, url: str) -> str:
    for src in (name, urlparse(url).path):
        if "." in src:
            return src.rsplit(".", 1)[-1].lower()[:32]
    return ""


def _fetch_bytes(url: str) -> bytes:
    """京东 OSS 对普通 requests 返回 403，需用 curl_cffi 伪装浏览器。"""
    from curl_cffi import requests as cffi_requests

    headers = {
        "Accept": "*/*",
        "Referer": "https://zcy-auction-view-pro.pf.jd.com/",
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/131.0.0.0 Safari/537.36"
        ),
    }
    kwargs = {
        "headers": headers,
        "timeout": 60,
        "impersonate": settings.jd_impersonate,
    }
    proxy = proxy_url()
    if proxy:
        kwargs["proxies"] = {"http": proxy, "https": proxy}
    resp = cffi_requests.get(url, **kwargs)
    if resp.status_code != 200:
        raise RuntimeError(f"HTTP {resp.status_code}")
    content = resp.content or b""
    if len(content) < 16 or content[:1] in (b"<", b"{") or content[:5].lower().startswith(b"<!doc"):
        raise RuntimeError(f"响应不是附件文件（{len(content)} bytes）")
    return content


def _sniff_ext(content: bytes, fallback: str) -> str:
    if content.startswith(b"%PDF"):
        return "pdf"
    if content.startswith(b"PK\x03\x04") and fallback in {"docx", "xlsx", "xls", "zip"}:
        return fallback or "zip"
    return fallback


def download_attachment(att: NoticeAttachment) -> None:
    url = ensure_https(att.file_url)
    content = _fetch_bytes(url)
    sha = hashlib.sha256(content).hexdigest()
    ext = _sniff_ext(content, _ext_from_name(att.file_name or "", url))
    folder = settings.attachment_dir / att.site_id / str(att.notice_raw_id)
    folder.mkdir(parents=True, exist_ok=True)
    base_name = att.file_name or f"{sha}.{ext or 'bin'}"
    if ext and not base_name.lower().endswith(f".{ext}"):
        base_name = f"{base_name}.{ext}"
    safe_name = re.sub(r"[\\\\/:*?\"<>|]+", "_", base_name)
    path = folder / safe_name
    path.write_bytes(content)
    att.local_path = str(path)
    att.file_ext = ext
    att.file_size = len(content)
    att.content_sha256 = sha
    att.download_status = "done"
    att.downloaded_at = now_shanghai()
    att.error_message = None
    if att.parse_status in ("parse_fail", "skipped"):
        att.parse_status = "pending"


def download_pending(*, limit: int = 50, site_id: Optional[str] = None) -> int:
    session = get_session()
    done = 0
    try:
        stmt = select(NoticeAttachment).where(
            NoticeAttachment.download_status.in_(["pending", "fail"])
        )
        if site_id:
            stmt = stmt.where(NoticeAttachment.site_id == site_id)
        stmt = stmt.order_by(
            case((NoticeAttachment.download_status == "pending", 0), else_=1),
            NoticeAttachment.id.asc(),
        ).limit(limit)
        rows = session.execute(stmt).scalars().all()
        for att in tqdm(rows, desc="download-att"):
            try:
                download_attachment(att)
                done += 1
            except Exception as exc:  # noqa: BLE001
                att.download_status = "fail"
                att.error_message = str(exc)[:1000]
            session.commit()
    finally:
        session.close()
    return done


def extract_text_from_file(path: Path) -> str:
    ext = path.suffix.lower().lstrip(".")
    with path.open("rb") as fh:
        head = fh.read(5)
    if ext == "pdf" or head == b"%PDF-":
        text = _pdf_text(path)
        if text.strip():
            return text
        ocr_text = _ocr_pdf(path)
        if ocr_text.strip():
            return ocr_text
        return ""
    if ext in {"png", "jpg", "jpeg", "bmp", "tif", "tiff", "webp"}:
        return _ocr_image(path)
    if ext in {"txt", "md", "csv", "json", "html", "htm"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext in {"docx"}:
        return _docx_text(path)
    if ext in {"xlsx", "xlsm"}:
        return _xlsx_text(path)
    return ""


def _pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


def _ocr_engine():
    try:
        from rapidocr_onnxruntime import RapidOCR
    except ImportError:
        return None
    return RapidOCR()


def _ocr_pil(engine, pil) -> str:
    import numpy as np

    if pil.mode != "RGB":
        pil = pil.convert("RGB")
    result, _elapse = engine(np.array(pil))
    if not result:
        return ""
    return "\n".join(str(row[1]) for row in result if len(row) > 1)


def _ocr_pdf(path: Path) -> str:
    """扫描件 PDF：抽出内嵌图片后走 RapidOCR。未安装时返回空。"""
    engine = _ocr_engine()
    if engine is None:
        return ""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        images = getattr(page, "images", None) or []
        for image in images:
            pil = getattr(image, "image", None)
            if pil is None:
                continue
            text = _ocr_pil(engine, pil)
            if text:
                parts.append(text)
    return "\n".join(parts).strip()


def _ocr_image(path: Path) -> str:
    engine = _ocr_engine()
    if engine is None:
        return ""
    from PIL import Image

    with Image.open(path) as pil:
        return _ocr_pil(engine, pil)


def _docx_text(path: Path) -> str:
    try:
        import docx
    except ImportError:
        return ""
    document = docx.Document(str(path))
    lines = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" ".join(cells))
    return "\n".join(lines).strip()


def _xlsx_text(path: Path) -> str:
    try:
        import openpyxl
    except ImportError:
        return ""
    book = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for sheet in book.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value).strip() for value in row]
                if any(cells):
                    lines.append(" ".join(cells))
    finally:
        book.close()
    return "\n".join(lines).strip()


def parse_attachment_row(session, att: NoticeAttachment, *, use_llm: bool = True) -> bool:
    """抽取附件文字，用本地模型解析标的（货物/服务），并在公告商品是占位时回写。"""
    if not att.local_path:
        att.parse_status = "skipped"
        att.error_message = "没有本地文件"
        return False
    path = Path(att.local_path)
    if not path.exists():
        att.parse_status = "parse_fail"
        att.error_message = "local file missing"
        return False
    # 已抽过正文则复用，避免大 PDF 重复解析拖慢 parse-att
    text = (att.extracted_text or "").strip()
    if not text:
        text = extract_text_from_file(path)
        att.extracted_text = text
    if not text.strip():
        att.parse_status = "skipped"
        att.error_message = "没有可抽取的文字（可能是扫描件）"
        return False
    if not use_llm:
        # 抓取阶段只下载：保留正文，解析留给 parse-att
        att.parse_status = "pending"
        att.error_message = None
        return False

    raw = session.get(NoticeRaw, att.notice_raw_id)
    title = (raw.title if raw else "") or att.file_name or ""
    model_parsed = extract_attachment_with_local_model(text, title=title)
    if not isinstance(model_parsed, dict):
        model_parsed = {}
    goods = model_parsed.get("goods") if isinstance(model_parsed.get("goods"), list) else []
    if goods_are_placeholder(goods):
        goods = []
    model_extra = model_parsed.get("extra") if isinstance(model_parsed.get("extra"), dict) else {}
    parsed = {
        "goods": goods,
        "extra": {
            "extractor": model_extra.get("extractor", "local_gguf_attachment"),
        },
    }
    # 京东联系人仍以详情 JSON 为准，模型结果只留底便于排查
    if is_jd_site(att.site_id):
        parsed["extra"]["local_model_raw"] = {
            k: model_parsed.get(k)
            for k in ("goods", "buyer_contact_name", "buyer_phone", "delivery_address")
        }
    else:
        for key in (
            "buyer_contact_name",
            "buyer_phone",
            "buyer_address",
            "buyer_contact",
            "delivery_address",
        ):
            value = model_parsed.get(key)
            if value not in (None, "", []):
                parsed[key] = value
    if not isinstance(parsed.get("extra"), dict):
        parsed["extra"] = {}
    parsed["extra"]["from_attachment_id"] = att.id
    parsed["extra"]["from_attachment"] = att.file_name
    if model_extra.get("extractor") in {"local_model_missing", "local_model_fail"}:
        att.parse_status = "parse_fail"
        att.error_message = (
            model_extra.get("local_model_error")
            or model_extra.get("model_missing")
            or "本地模型不可用"
        )[:1000]
        att.parsed_json = dumps(parsed)
        return False
    # JSON 失败：标 parse_fail 方便重跑，不要当成已解析成功
    if str(model_extra.get("extractor") or "").endswith("+json_fail"):
        att.parsed_json = dumps(parsed)
        att.parse_status = "parse_fail"
        att.error_message = "模型输出不是合法 JSON（可重跑 parse-att）"[:1000]
        return False
    name = att.file_name or ""
    is_score_doc = any(token in name for token in ("评分", "办法", "须知"))
    # 「技术标准/参数」可能含真实标的，不因「标准」二字跳过；纯评分办法跳过
    if is_score_doc and "参数" not in name and "报价" not in name:
        # 模型偶发从评分表里幻觉出一行，不入库
        parsed["goods"] = []
        att.parsed_json = dumps(parsed)
        att.parse_status = "skipped"
        att.error_message = "评分/标准类附件无采购标的"
        return False
    if is_jd_site(att.site_id) and not parsed.get("goods"):
        att.parsed_json = dumps(parsed)
        att.parse_status = "parse_fail"
        att.error_message = "未抽出产品清单（可重跑 parse-att）"
        return False
    att.parsed_json = dumps(parsed)
    att.parse_status = "parsed"
    att.parsed_at = now_shanghai()
    att.error_message = None
    if raw is not None and (
        parsed.get("goods") or (not is_jd_site(att.site_id))
    ):
        _merge_into_notice(session, raw, parsed)
    return True


def _merge_into_notice(session, raw: NoticeRaw, parsed: dict) -> None:
    """附件结果补进公告。摘要 SKU / 详见附件时用附件清单覆盖。"""
    detail = raw.detail_json if isinstance(raw.detail_json, dict) else {}
    if raw.parse_status in ("pending", "parse_fail", "skipped"):
        # 先按详情补字段，再用附件清单覆盖摘要商品
        merged = enrich_from_raw(dict(parsed), raw)
        if should_replace_goods_with_attachment(
            merged.get("goods"), parsed.get("goods"), detail=detail
        ):
            merged["goods"] = parsed["goods"]
        elif parsed.get("goods") and not merged.get("goods"):
            merged["goods"] = parsed["goods"]
        save_parsed(session, raw, merged)
        return
    existing = session.execute(
        select(NoticeParsed).where(NoticeParsed.notice_raw_id == raw.id)
    ).scalar_one_or_none()
    goods_rows = list(
        session.execute(
            select(NoticeGoods).where(NoticeGoods.notice_raw_id == raw.id)
        ).scalars().all()
    )
    base = parsed_row_to_dict(existing, goods_rows) if existing else {}
    for key, value in parsed.items():
        if key == "extra":
            extra = base.get("extra") if isinstance(base.get("extra"), dict) else {}
            if isinstance(value, dict):
                extra = {**extra, **value}
            base["extra"] = extra
            continue
        if key == "goods":
            if should_replace_goods_with_attachment(
                base.get("goods"), value, detail=detail
            ):
                base["goods"] = value
            continue
        if base.get(key) in (None, "", []) and value not in (None, "", []):
            base[key] = value
    base = enrich_from_raw(base, raw)
    # enrich 可能又从详情 JSON 填回摘要 SKU，附件清单优先
    if should_replace_goods_with_attachment(
        base.get("goods"), parsed.get("goods"), detail=detail
    ):
        base["goods"] = parsed["goods"]
    save_parsed(session, raw, base)


def parse_attachments(*, limit: int = 50, site_id: Optional[str] = None, use_llm: bool = True) -> int:
    session = get_session()
    ok = 0
    try:
        stmt = select(NoticeAttachment).where(
            NoticeAttachment.download_status == "done",
            NoticeAttachment.parse_status.in_(["pending", "parse_fail"]),
        )
        if site_id:
            stmt = stmt.where(NoticeAttachment.site_id == site_id)
        stmt = stmt.order_by(NoticeAttachment.id.asc()).limit(limit)
        rows = session.execute(stmt).scalars().all()
        for att in tqdm(rows, desc="parse-att"):
            try:
                if parse_attachment_row(session, att, use_llm=use_llm):
                    ok += 1
            except Exception as exc:  # noqa: BLE001
                att.parse_status = "parse_fail"
                att.error_message = str(exc)[:1000]
            session.commit()
    finally:
        session.close()
    return ok


def register_attachments(session, raw: NoticeRaw, files: list) -> list:
    """把详情里的附件登记到 notice_attachment，已有 URL 不重复插入。"""
    rows = []
    for item in files:
        url = ensure_https((item or {}).get("file_url") or "")
        if not url:
            continue
        exists = session.execute(
            select(NoticeAttachment).where(
                NoticeAttachment.notice_raw_id == raw.id,
                NoticeAttachment.file_url == url,
            )
        ).scalar_one_or_none()
        if exists:
            rows.append(exists)
            continue
        row = NoticeAttachment(
            notice_raw_id=raw.id,
            site_id=raw.site_id,
            notice_key=raw.notice_key,
            file_name=item.get("file_name"),
            file_url=url,
            download_status="pending",
            parse_status="pending",
        )
        session.add(row)
        rows.append(row)
    session.flush()
    return rows


def ingest_attachments(session, raw: NoticeRaw, files: list, *, use_llm: bool = False) -> int:
    """登记并下载附件。默认不跑模型（抓取快路径）；解析用 parse-att。"""
    ok = 0
    for att in register_attachments(session, raw, files):
        if att.download_status != "done" or not att.local_path or not Path(att.local_path).exists():
            try:
                download_attachment(att)
                ok += 1
            except Exception as exc:  # noqa: BLE001
                att.download_status = "fail"
                att.error_message = str(exc)[:1000]
                continue
        if use_llm and att.parse_status != "parsed":
            try:
                if parse_attachment_row(session, att, use_llm=True):
                    ok += 1
            except Exception as exc:  # noqa: BLE001
                att.parse_status = "parse_fail"
                att.error_message = str(exc)[:1000]
        elif att.download_status == "done" and att.parse_status not in ("parsed", "skipped"):
            att.parse_status = "pending"
    return ok


def sync_missing_jd_attachments(*, limit: int = 100, use_llm: bool = False) -> int:
    """历史京东公告：详情里已有附件列表、但还没入库的，补下载（默认不解析）。"""
    from sqlalchemy import text

    session = get_session()
    client = JdQuanClient()
    done = 0
    try:
        id_rows = session.execute(
            text(
                """
                SELECT r.id
                FROM notice_raw r
                WHERE r.site_id IN ('jd-quan-yicai', 'jd-qing-huicai')
                  AND JSON_TYPE(JSON_EXTRACT(r.detail_json, '$.attachmentInfoList')) = 'ARRAY'
                  AND JSON_LENGTH(r.detail_json, '$.attachmentInfoList') > 0
                  AND NOT EXISTS (
                    SELECT 1 FROM notice_attachment a WHERE a.notice_raw_id = r.id
                  )
                ORDER BY r.id DESC
                LIMIT :limit
                """
            ),
            {"limit": limit},
        ).fetchall()
        for (raw_id,) in id_rows:
            raw = session.get(NoticeRaw, raw_id)
            if raw is None or not isinstance(raw.detail_json, dict):
                continue
            files = client.extract_attachments(raw.detail_json)
            if not files:
                continue
            done += ingest_attachments(session, raw, files, use_llm=use_llm)
            session.commit()
    finally:
        session.close()
    return done
